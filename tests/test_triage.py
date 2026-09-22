from __future__ import annotations

import unittest
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from api.enquiries import EnquiryRepository
from scripts.stage_06_triage.routing_policy import apply_routing_policy
from scripts.stage_06_triage.schemas import (
    Category, EnquiryClassification, RouteTo, TriageAction,
)
from scripts.stage_06_triage.sensitive_rules import detect_sensitive_terms


def classification(**overrides):
    values = dict(
        summary="A routine expense question",
        category=Category.FINANCE_OPERATIONS,
        subcategory="expense_claim",
        is_sensitive=False,
        action=TriageAction.DRAFT_RESPONSE,
        route_to=RouteTo.FINANCE_TEAM,
        missing_info=(),
    )
    values.update(overrides)
    return EnquiryClassification(**values)


class TriageTests(unittest.TestCase):
    def test_sensitive_rule_overrides_safe_model_classification(self):
        flags = detect_sensitive_terms(
            "I need to change my address because someone is stalking me."
        )
        decision = apply_routing_policy(
            classification(category=Category.STUDENT_ADMIN), flags
        )
        self.assertFalse(decision.allow_generation)
        self.assertEqual(decision.action, TriageAction.SPECIALIST_REFERRAL)
        self.assertEqual(decision.route_to, RouteTo.HARASSMENT_AND_MISCONDUCT)

    def test_llm_sensitive_flag_blocks_generation_without_keyword_rule(self):
        decision = apply_routing_policy(classification(
            category=Category.UNSUPPORTED,
            is_sensitive=True,
            route_to=RouteTo.STUDENT_COUNSELLING_AND_WELLBEING,
        ))
        self.assertFalse(decision.allow_generation)
        self.assertEqual(decision.action, TriageAction.SPECIALIST_REFERRAL)

    def test_missing_information_requests_clarification(self):
        decision = apply_routing_policy(classification(missing_info=("module_code",)))
        self.assertEqual(decision.action, TriageAction.REQUEST_CLARIFICATION)
        self.assertFalse(decision.allow_generation)

    def test_classifier_clarification_action_is_enforced_by_policy(self):
        decision = apply_routing_policy(classification(
            action=TriageAction.REQUEST_CLARIFICATION,
        ))
        self.assertEqual(decision.action, TriageAction.REQUEST_CLARIFICATION)
        self.assertFalse(decision.allow_generation)

    def test_classifier_manual_review_action_is_enforced_by_policy(self):
        decision = apply_routing_policy(classification(
            action=TriageAction.MANUAL_REVIEW,
        ))
        self.assertEqual(decision.action, TriageAction.MANUAL_REVIEW)
        self.assertFalse(decision.allow_generation)

    def test_unsupported_category_routes_to_manual_review(self):
        decision = apply_routing_policy(classification(
            category=Category.UNSUPPORTED,
            subcategory="facilities_safety",
            action=TriageAction.MANUAL_REVIEW,
            route_to=RouteTo.MANUAL_TRIAGE,
        ))
        self.assertEqual(decision.action, TriageAction.MANUAL_REVIEW)
        self.assertEqual(decision.route_to, RouteTo.MANUAL_TRIAGE)
        self.assertFalse(decision.allow_generation)
        self.assertFalse(decision.classification.is_sensitive)

    def test_unsupported_takes_precedence_over_missing_information(self):
        decision = apply_routing_policy(classification(
            category=Category.UNSUPPORTED,
            action=TriageAction.REQUEST_CLARIFICATION,
            missing_info=("fault severity",),
        ))
        self.assertEqual(decision.action, TriageAction.MANUAL_REVIEW)
        self.assertEqual(decision.route_to, RouteTo.MANUAL_TRIAGE)
        self.assertFalse(decision.allow_generation)

    def test_answerable_category_allows_staff_reviewed_draft(self):
        decision = apply_routing_policy(classification())
        self.assertTrue(decision.allow_generation)
        self.assertEqual(decision.action, TriageAction.DRAFT_RESPONSE)

    def test_repository_redacts_sensitive_text_and_exports_excel(self):
        repository = EnquiryRepository(database_url=None)
        record = {
            "enquiry_id":"id-1", "created_at_utc":"2026-09-20T12:00:00Z",
            "enquiry":"Sensitive details", "summary":"Needs support",
            "category":"unsupported", "subcategory":"wellbeing",
            "is_sensitive":True, "action":"specialist_referral",
            "route_to":"student_counselling_and_wellbeing", "missing_info":[],
            "answer_status":None, "draft_response":None, "citation_ids":[],
            "review_status":"pending",
        }
        repository.save(record)
        self.assertEqual(repository.list()[0]["enquiry"], "[Sensitive enquiry redacted]")
        workbook = load_workbook(BytesIO(repository.export_xlsx()))
        self.assertEqual(workbook.active["A1"].value, "enquiry_id")
        self.assertEqual(workbook.active["C2"].value, "[Sensitive enquiry redacted]")

    def test_excel_export_serializes_timezone_aware_database_timestamp(self):
        repository = EnquiryRepository(database_url=None)
        repository._records.append({
            "enquiry_id": "id-2",
            "created_at_utc": datetime(2026, 9, 20, 5, 0, tzinfo=timezone.utc),
            "enquiry": "Synthetic enquiry",
            "summary": "Synthetic",
            "category": "finance_operations",
            "subcategory": "expense_claim",
            "is_sensitive": False,
            "action": "draft_response",
            "route_to": "finance_team",
            "missing_info": [],
            "answer_status": "answered",
            "draft_response": "Draft",
            "citation_ids": [],
            "review_status": "pending",
            "staff_object_id": "staff-1",
            "staff_display_name": "Staff User",
        })
        workbook = load_workbook(BytesIO(repository.export_xlsx()))
        self.assertEqual(workbook.active["B2"].value, "2026-09-20T05:00:00+00:00")


if __name__ == "__main__":
    unittest.main()
